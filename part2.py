"""Part 2. Stock-level cleaning and three book-to-price portfolios.

    python part2.py
"""

from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
STOCKS = ROOT / "data" / "stocks.csv"
FF = ROOT / "data" / "ff_msci.xlsx"
OUT = ROOT / "exhibits" / "part2"
FORM_START = pd.Period("1998-12", freq="M")
FORM_END = pd.Period("2018-11", freq="M")

plt.rcParams.update({
    "figure.dpi": 140,
    "savefig.bbox": "tight",
    "savefig.facecolor": "white",
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "axes.edgecolor": "#333333",
    "axes.labelcolor": "#222222",
    "axes.titlecolor": "#1F4E79",
    "text.color": "#222222",
    "axes.grid": True,
    "grid.color": "#E6E6E6",
    "grid.linewidth": 0.7,
    "axes.axisbelow": True,
    "font.size": 10,
})
BLUE, ORANGE = "#4C78A8", "#F58518"


def month_stamp(value):
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        return pd.Period(pd.Timestamp(value), freq="M")
    except (ValueError, TypeError):
        return None


def as_period(value):
    if value is None or isinstance(value, datetime):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) == 6 and text.isdigit():
        return pd.Period(year=int(text[:4]), month=int(text[4:]), freq="M")
    return None


def save_fig(fig, name):
    fig.tight_layout()
    fig.savefig(OUT / name)
    plt.close(fig)


def load_hml():
    wb = load_workbook(FF, read_only=True, data_only=True)
    rows = list(wb["FF - 3 factor returns"].iter_rows(values_only=True))
    wb.close()
    header = next(i for i, row in enumerate(rows) if row and str(row[1]).strip() == "Mkt-RF")
    hml_at = next(i for i, cell in enumerate(rows[header]) if cell is not None and str(cell).strip() == "HML")
    values = {}
    for row in rows[header + 1:]:
        period = as_period(row[0])
        if period is None:
            break
        values[period] = float(row[hml_at]) / 100.0
    return pd.Series(values).sort_index()


def load_usa():
    wb = load_workbook(FF, read_only=True, data_only=True)
    rows = list(wb["MSCI - factor totret index"].iter_rows(values_only=True))
    wb.close()
    header_at = next(i for i, row in enumerate(rows) if row and str(row[0]).strip() == "Date")
    usa_at = next(i for i, cell in enumerate(rows[header_at]) if cell and "USA" in str(cell).upper() and "VALUE" not in str(cell).upper())
    levels = {}
    for row in rows[header_at + 1:]:
        period = month_stamp(row[0] if row else None)
        if period is None:
            if levels:
                break
            continue
        levels[period] = float(str(row[usa_at]).replace(",", ""))
    return pd.Series(levels).sort_index().pct_change().dropna()


def load_stocks():
    df = pd.read_csv(STOCKS, parse_dates=["date"])
    df["form"] = df["date"].dt.to_period("M")
    df["bp"] = 1.0 - df["Pb"]
    df["r"] = df.groupby("form")["R1M_Usd"].transform(lambda s: s.clip(upper=s.quantile(0.99)))
    return df


def write_cleaning(df):
    counts = df.groupby("form").size()
    coverage = df.groupby("form")["R1M_Usd"].apply(lambda s: s.notna().mean())
    caps = df.groupby("form")["R1M_Usd"].quantile(0.99)
    months = counts.index.to_timestamp("M")

    fig, ax = plt.subplots(figsize=(7.4, 3.6))
    ax.plot(months, counts.to_numpy(), color=BLUE, lw=1.4)
    ax.set_ylabel("Number of stocks")
    ax.set_title("Stocks in the file each month")
    save_fig(fig, "2a_counts.png")

    raw = df["R1M_Usd"]
    fig, ax = plt.subplots(figsize=(7.4, 3.6))
    ax.hist(raw[raw <= 1.5], bins=60, color=BLUE, edgecolor="white", linewidth=0.3)
    ax.axvline(caps.median(), color=ORANGE, lw=1.4, zorder=3)
    ax.set_xlim(-1.0, 1.5)
    ax.set_xlabel("Forward one-month return")
    ax.set_ylabel("Stock-months")
    n_above = int((raw > 1.5).sum())
    ax.set_title(f"Return distribution ({n_above} stock-months above 1.5; max {raw.max():.2f})")
    save_fig(fig, "2a_returns.png")

    fig, ax = plt.subplots(figsize=(7.4, 3.6))
    ax.plot(months, coverage.to_numpy(), color=BLUE, lw=1.4)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("Share non-missing")
    ax.set_title("Return coverage each month")
    save_fig(fig, "2a_coverage.png")

    fig, ax = plt.subplots(figsize=(7.4, 3.6))
    ax.hist(df["Pb"], bins=np.arange(0.005, 1.015, 0.01), color=BLUE, edgecolor="white", linewidth=0.2)
    ax.axvline(0.5, color=ORANGE, lw=1.4, zorder=3)
    ax.set_xlabel("Price-to-book percentile (1 = highest P/B)")
    ax.set_ylabel("Stock-months")
    ax.set_title("Standardized price-to-book")
    save_fig(fig, "2a_pb.png")

    rhos = []
    for _, g in df.groupby("form"):
        left = g["Pb"].rank()
        right = (g["Mkt_Cap_3M_Usd"] - g["Bv"]).rank()
        rhos.append(left.corr(right))
    rhos = pd.Series(rhos)
    text = "\n".join([
        "Exhibit 2A — does the price-to-book percentile behave like market value / book?",
        "Within each month, Spearman correlation of Pb with (market-cap percentile − book percentile).",
        "The percentiles are not divided. Pb is price-to-book, so a positive correlation is the expected sign.",
        f"Sample: {df['form'].min()} to {df['form'].max()}, {df['form'].nunique()} months, {df['stock_id'].nunique()} stocks.",
        "",
        f"Median Spearman correlation: {rhos.median():.3f}",
        f"Share of months with a positive correlation: {float((rhos > 0).mean()):.3f}",
    ])
    (OUT / "2a_pb_check.txt").write_text(text + "\n")


def portfolio(g, kind):
    g = g.dropna(subset=["r", "bp", "Mkt_Cap_3M_Usd"])
    n = len(g)
    if kind == "mkt":
        raw = g["Mkt_Cap_3M_Usd"].to_numpy()
    elif kind == "ew":
        k = max(1, int(np.floor(len(g) * 0.2 + 0.5)))
        order = np.lexsort((g["stock_id"].to_numpy(), -g["bp"].to_numpy()))
        raw = np.zeros(n)
        raw[order[:k]] = 1.0
    elif kind == "score":
        raw = g["bp"].to_numpy()
    else:
        raw = (g["bp"] * g["Mkt_Cap_3M_Usd"]).to_numpy()
    weights = raw / raw.sum()
    ret = float(weights @ g["r"].to_numpy())
    held = int((weights > 0).sum())
    return ret, held, float(weights.max())


def run_portfolios(df):
    use = df[df["form"].between(FORM_START, FORM_END)]
    out = {name: {} for name in ("mkt", "ew", "score", "score_cap")}
    last = {}
    for form, g in use.groupby("form"):
        ret_month = form + 1
        for name in out:
            ret, held, max_w = portfolio(g, name)
            out[name][ret_month] = ret
            if form == FORM_END:
                last[name] = (held, max_w)
    returns = pd.DataFrame({name: pd.Series(values) for name, values in out.items()}).sort_index()
    expect = (pd.Period("2018-12") - pd.Period("1999-01")).n + 1
    if len(returns) != expect or returns.index.min() != pd.Period("1999-01") or returns.index.max() != pd.Period("2018-12"):
        raise RuntimeError(f"portfolio window {returns.index.min()}–{returns.index.max()} n={len(returns)}")
    return returns, last


def ann_geo(series):
    return float(np.prod(1.0 + series.to_numpy()) ** (12.0 / len(series)) - 1.0)


def ann_vol(series):
    return float(series.std(ddof=1) * np.sqrt(12))


def write_results(returns, last, usa, hml):
    usa = usa.reindex(returns.index)
    hml = hml.reindex(returns.index)
    if usa.isna().any() or hml.isna().any():
        raise RuntimeError("MSCI USA or HML does not cover 1999–2018")
    labels = {
        "mkt": "Cap-weight proxy",
        "ew": "EW top 20% B/P",
        "score": "B/P score",
        "score_cap": "B/P score × cap",
    }
    lines = [
        "Exhibit 2C — book-to-price portfolios, formation Dec 1998–Nov 2018 (return months Jan 1999–Dec 2018)",
        "Long-only, fully invested. Book-to-price score = 1 − price-to-book percentile. Returns winsorized within each month at the 99th percentile.",
        "ann_return is the compounded annualized return. ann_vol, tracking error, and max_w are in percent.",
        "ann_vol and tracking error are monthly std × sqrt(12). n_names is a count. Tracking error is versus the cap-weight proxy.",
        "Names and max weight are measured on the Nov 2018 formation date.",
        "",
        f"{'strategy':<20}{'ann_return':>12}{'ann_vol':>10}{'te':>8}{'n_names':>9}{'max_w':>8}",
    ]
    for key, label in labels.items():
        excess = returns[key] - returns["mkt"]
        lines.append(
            f"{label:<20}{100 * ann_geo(returns[key]):12.2f}{100 * ann_vol(returns[key]):10.2f}"
            f"{100 * ann_vol(excess):8.2f}{last[key][0]:9d}{100 * last[key][1]:8.2f}"
        )
    (OUT / "2c_value.txt").write_text("\n".join(lines) + "\n")

    corr = float(returns["mkt"].corr(usa))
    text = "\n".join([
        "Exhibit 2B — cap-weight proxy versus MSCI USA",
        "Proxy weights are proportional to the standardized market-cap percentile. Return months Jan 1999–Dec 2018.",
        "MSCI USA is the gross large+mid-cap total-return index. Both series are monthly total returns.",
        "",
        f"Months: {len(returns)}",
        f"Correlation: {corr:.3f}",
    ])
    (OUT / "2b_market.txt").write_text(text + "\n")

    fig, ax = plt.subplots(figsize=(5.6, 5.4))
    ax.scatter(100 * usa, 100 * returns["mkt"], s=22, c=BLUE, alpha=0.55, zorder=2)
    lo = min(100 * usa.min(), 100 * returns["mkt"].min())
    hi = max(100 * usa.max(), 100 * returns["mkt"].max())
    pad = 0.06 * (hi - lo)
    ax.plot([lo, hi], [lo, hi], color=ORANGE, lw=1.8, zorder=4, label="Same as market")
    ax.set_xlim(lo - pad, hi + pad)
    ax.set_ylim(lo - pad, hi + pad)
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("MSCI USA (% per month)")
    ax.set_ylabel("Cap-weight proxy (% per month)")
    ax.set_title(f"Cap-weight proxy vs MSCI USA (corr {corr:.2f})", pad=8)
    ax.legend(frameon=False, loc="upper left")
    save_fig(fig, "2b_market.png")

    excess = pd.DataFrame({
        "EW top 20%": returns["ew"] - returns["mkt"],
        "B/P score": returns["score"] - returns["mkt"],
        "Score × cap": returns["score_cap"] - returns["mkt"],
        "HML": hml,
    })
    c = excess.corr()
    lines = [
        "Exhibit 2D — correlation of value-portfolio excess returns",
        "Excess return is the portfolio minus the cap-weight proxy. HML is the Ken French value spread over the same months (Jan 1999–Dec 2018).",
        "",
        f"{'':<14}" + "".join(f"{name:>14}" for name in c.columns),
    ]
    for name, row in c.iterrows():
        lines.append(f"{name:<14}" + "".join(f"{value:14.2f}" for value in row))
    (OUT / "2d_corr.txt").write_text("\n".join(lines) + "\n")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    df = load_stocks()
    write_cleaning(df)
    returns, last = run_portfolios(df)
    write_results(returns, last, load_usa(), load_hml())


if __name__ == "__main__":
    main()
