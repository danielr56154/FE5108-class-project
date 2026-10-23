"""Part 1. Factor premia, MSCI indexes, construction correlations, momentum vs volatility.

    python part1.py
"""

from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import linregress

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "data" / "ff_msci.xlsx"
OUT = ROOT / "exhibits"

plt.rcParams.update({
    "figure.dpi": 140,
    "savefig.bbox": "tight",
    "axes.spines.top": False,
    "axes.spines.right": False,
    "font.size": 10,
})


def month_stamp(value):
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        return pd.Period(pd.Timestamp(value), freq="M")
    except (ValueError, TypeError):
        return None


def as_period(value):
    if value is None or isinstance(value, datetime):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) == 6 and text.isdigit():
        return pd.Period(year=int(text[:4]), month=int(text[4:]), freq="M")
    return None


def sheet_rows(name):
    wb = load_workbook(SRC, read_only=True, data_only=True)
    rows = [list(row) for row in wb[name].iter_rows(values_only=True)]
    wb.close()
    return rows


def column(header, label):
    for i, cell in enumerate(header):
        if cell is not None and str(cell).strip() == label:
            return i
    raise RuntimeError(f"missing column {label}")


def french_block(rows, header_at, columns):
    """Monthly percent returns from the first Ken French block under a header row."""
    indexes = {name: column(rows[header_at], name) for name in columns}
    data = {name: {} for name in columns}
    for row in rows[header_at + 1:]:
        period = as_period(row[0])
        if period is None:
            break
        for name, i in indexes.items():
            value = row[i]
            data[name][period] = np.nan if value is None or float(value) <= -99 else float(value)
    return pd.DataFrame(data).sort_index()


def load_factors():
    rows = sheet_rows("FF - 3 factor returns")
    header = next(i for i, row in enumerate(rows) if row and str(row[1]).strip() == "Mkt-RF")
    raw = french_block(rows, header, ["Mkt-RF", "SMB", "HML", "RF"])
    return raw / 100.0


def load_momentum():
    rows = sheet_rows("FF - portfolios sorted on pmom")
    header = next(i for i, row in enumerate(rows) if any(str(c).strip() == "Lo PRIOR" for c in row if c is not None))
    raw = french_block(rows, header, ["Lo PRIOR", "Hi PRIOR"])
    return (raw["Hi PRIOR"] - raw["Lo PRIOR"]) / 100.0


def load_low_beta():
    rows = sheet_rows("FF - portfolios sorted on beta")
    header = next(i for i, row in enumerate(rows) if any(str(c).strip() == "Lo 10" for c in row if c is not None))
    raw = french_block(rows, header, ["Lo 10", "Hi 10"])
    return (raw["Lo 10"] - raw["Hi 10"]) / 100.0


def load_msci():
    rows = sheet_rows("MSCI - factor totret index")
    header_at = next(i for i, row in enumerate(rows) if row and str(row[0]).strip() == "Date")
    header = [str(c).strip() if c is not None else "" for c in rows[header_at]]
    keys = {}
    for i, name in enumerate(header):
        upper = name.upper()
        if "VALUE" in upper:
            keys[i] = "value"
        elif "MINIMUM" in upper:
            keys[i] = "minvol"
        elif "EQUAL" in upper:
            keys[i] = "ew"
        elif "MOMENTUM" in upper:
            keys[i] = "mom"
        elif "USA" in upper:
            keys[i] = "usa"
    levels = {name: {} for name in keys.values()}
    for row in rows[header_at + 1:]:
        period = month_stamp(row[0] if row else None)
        if period is None:
            if levels["usa"]:
                break
            continue
        for i, name in keys.items():
            levels[name][period] = float(str(row[i]).replace(",", ""))
    level = pd.DataFrame(levels).sort_index()
    return level.pct_change().dropna()


def window(frame, start, end):
    out = frame.loc[start:end]
    expect = (pd.Period(end, freq="M") - pd.Period(start, freq="M")).n + 1
    if len(out) != expect or out.index.min() != pd.Period(start) or out.index.max() != pd.Period(end):
        raise RuntimeError(f"{start} to {end}: got {out.index.min()}–{out.index.max()} (n={len(out)})")
    return out


def mean_t(series):
    n = len(series)
    mu = float(series.mean())
    t = mu / (float(series.std(ddof=1)) / np.sqrt(n))
    return mu, t, n


def ann_stats(series, rf):
    excess = series - rf
    vol = float(series.std(ddof=1)) * np.sqrt(12)
    sharpe = float(excess.mean()) / float(series.std(ddof=1)) * np.sqrt(12)
    return 12 * float(series.mean()), vol, sharpe


def save_fig(fig, name, note):
    fig.text(0.0, -0.02, note, fontsize=7, color="#444444", ha="left", transform=fig.transFigure)
    fig.tight_layout()
    fig.savefig(OUT / name)
    plt.close(fig)


def write_premia(factors):
    lines = [
        "Exhibit 1A — arithmetic mean and t-stat of monthly factor premia",
        "Ken French US factors, percent per month. t = mean / (std / sqrt(n)). No autocorrelation adjustment.",
        "Pre-publication: Jan 1927–Dec 1992. Post-publication: Jan 1993–Jul 2023.",
        "",
        f"{'window':<28}{'factor':<10}{'n':>6}{'mean':>10}{'t':>8}",
    ]
    for label, start, end in (
        ("Jan 1927–Dec 1992", "1927-01", "1992-12"),
        ("Jan 1993–Jul 2023", "1993-01", "2023-07"),
    ):
        block = window(factors, start, end)
        for name in ("Mkt-RF", "SMB", "HML"):
            mu, t, n = mean_t(block[name])
            lines.append(f"{label:<28}{name:<10}{n:6d}{100 * mu:10.3f}{t:8.2f}")
    (OUT / "1a_premia.txt").write_text("\n".join(lines) + "\n")


def write_performance(msci, rf):
    sample = window(msci, "1995-01", "2023-07")
    rf = window(rf, "1995-01", "2023-07")
    series = {
        "USA": sample["usa"],
        "Value": sample["value"],
        "Min vol": sample["minvol"],
        "Equal weight": sample["ew"],
        "Momentum": sample["mom"],
        "Multi-factor": sample[["value", "minvol", "ew", "mom"]].mean(axis=1),
        "T-bill": rf,
    }
    rows = {name: ann_stats(values, rf) for name, values in series.items()}
    lines = [
        "Exhibit 1C — annualized average return, volatility, and Sharpe ratio",
        "MSCI USA gross total-return indexes, USD, Jan 1995–Jul 2023. T-bill is Ken French RF.",
        "Multi-factor is the equal-weighted average of Value, Min vol, Equal weight, and Momentum each month.",
        "Annualized average return = 12 × arithmetic monthly mean. Sharpe = (ann. return − ann. RF) / (monthly std × sqrt(12)).",
        "",
        f"{'strategy':<16}{'ann_return':>12}{'ann_vol':>10}{'sharpe':>8}",
    ]
    for name, (ret, vol, sharpe) in rows.items():
        lines.append(f"{name:<16}{100 * ret:12.2f}{100 * vol:10.2f}{sharpe:8.2f}")
    (OUT / "1c_performance.txt").write_text("\n".join(lines) + "\n")

    order = sorted(rows, key=lambda name: rows[name][2])
    colors = ["#4C78A8" if name != "Multi-factor" else "#F58518" for name in order]
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 4.4), sharey=True)
    y = np.arange(len(order))
    axes[0].barh(y, [100 * rows[name][0] for name in order], color=colors)
    axes[1].barh(y, [rows[name][2] for name in order], color=colors)
    axes[0].set_xlabel("Annualized average return (%)")
    axes[1].set_xlabel("Sharpe ratio")
    axes[0].set_yticks(y, order)
    axes[1].axvline(0, color="#888888", lw=0.6)
    fig.suptitle("MSCI USA factor indexes, Jan 1995–Jul 2023", y=1.02)
    save_fig(fig, "1c_performance.png",
             "Gross total return, USD. Multi-factor (orange) averages the four smart-beta indexes. Sorted by Sharpe. Source: MSCI, Ken French RF.")


def write_corr(msci, factors, mom, low_beta):
    sample = window(msci, "1995-01", "2023-07")
    usa = sample["usa"]
    spreads = pd.DataFrame({
        "Value − USA": sample["value"] - usa,
        "Min vol − USA": sample["minvol"] - usa,
        "EW − USA": sample["ew"] - usa,
        "Mom − USA": sample["mom"] - usa,
    })
    french = pd.DataFrame({
        "HML": window(factors["HML"], "1995-01", "2023-07"),
        "SMB": window(factors["SMB"], "1995-01", "2023-07"),
        "FF momentum": window(mom, "1995-01", "2023-07"),
        "Low − high beta": window(low_beta, "1995-01", "2023-07"),
    })
    corr = pd.concat([spreads, french], axis=1).corr().loc[spreads.columns, french.columns]
    lines = [
        "Exhibit 1D — correlation of MSCI factor spreads with Fama-French spreads",
        "Jan 1995–Jul 2023. MSCI spreads are the smart-beta index minus MSCI USA.",
        "FF momentum is value-weighted Hi PRIOR minus Lo PRIOR. Low − high beta is value-weighted Lo 10 minus Hi 10.",
        "",
        f"{'':<16}" + "".join(f"{name:>16}" for name in corr.columns),
    ]
    for name, row in corr.iterrows():
        lines.append(f"{name:<16}" + "".join(f"{value:16.2f}" for value in row))
    (OUT / "1d_corr.txt").write_text("\n".join(lines) + "\n")

    fig, ax = plt.subplots(figsize=(7.2, 3.8))
    image = ax.imshow(corr.to_numpy(), cmap="coolwarm", vmin=-1, vmax=1)
    ax.set_xticks(range(corr.shape[1]), corr.columns, rotation=20, ha="right")
    ax.set_yticks(range(corr.shape[0]), corr.index)
    for i in range(corr.shape[0]):
        for j in range(corr.shape[1]):
            ax.text(j, i, f"{corr.iloc[i, j]:.2f}", ha="center", va="center", fontsize=9)
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.04)
    ax.set_title("Factor-spread correlations, Jan 1995–Jul 2023")
    save_fig(fig, "1d_corr.png",
             "MSCI gross total return minus MSCI USA, vs Ken French value-weighted spreads. USA large and mid cap.")


def write_momentum_vol(mkt, mom):
    vol = mkt.rolling(12).std(ddof=1) * np.sqrt(12)
    paired = pd.DataFrame({"mom": mom, "vol": vol.shift(1)}).dropna()
    paired = window(paired, "1995-12", "2023-07")
    med = float(paired["vol"].median())
    low = paired.loc[paired["vol"] < med, "mom"]
    high = paired.loc[paired["vol"] >= med, "mom"]
    fit = linregress(100 * paired["vol"], 100 * paired["mom"])
    lines = [
        "Exhibit 1E — momentum spread and lagged 12-month market volatility",
        "Dec 1995–Jul 2023. Momentum is the value-weighted high-minus-low prior-return decile, percent per month.",
        "Lagged vol is the annualized standard deviation of Mkt-RF over the 12 months ending the previous month.",
        "High vol means lagged vol at or above the sample median. t-stats ignore autocorrelation.",
        "",
        f"Median lagged ann. vol: {100 * med:.2f}%",
        f"{'bucket':<12}{'n':>6}{'mean':>10}{'t':>8}",
    ]
    for label, series in (("Low vol", low), ("High vol", high), ("All", paired["mom"])):
        mu, t, n = mean_t(series)
        lines.append(f"{label:<12}{n:6d}{100 * mu:10.3f}{t:8.2f}")
    lines += [
        "",
        "OLS: monthly momentum spread (%) = a + b × lagged ann. vol (%)",
        f"a = {fit.intercept:.3f}    b = {fit.slope:.4f}    t(b) = {fit.slope / fit.stderr:.2f}    n = {len(paired)}    R = {fit.rvalue:.3f}",
    ]
    (OUT / "1e_regimes.txt").write_text("\n".join(lines) + "\n")

    fig, ax = plt.subplots(figsize=(6.6, 4.4))
    ax.scatter(100 * paired["vol"], 100 * paired["mom"], s=12, alpha=0.45, color="#4C78A8", linewidths=0)
    x = np.linspace(100 * paired["vol"].min(), 100 * paired["vol"].max(), 50)
    ax.plot(x, fit.intercept + fit.slope * x, color="#E45756", lw=1.5)
    ax.axhline(0, color="#888888", lw=0.6)
    ax.set_xlabel("Lagged 12-month market volatility (% ann.)")
    ax.set_ylabel("Momentum spread (% per month)")
    ax.set_title("Momentum vs lagged market volatility, Dec 1995–Jul 2023")
    save_fig(fig, "1e_momentum_vol.png",
             "Value-weighted FF momentum decile spread. Volatility uses Mkt-RF through the previous month. Ken French.")


def main():
    OUT.mkdir(exist_ok=True)
    factors = load_factors()
    mom = load_momentum().rename("mom")
    low_beta = load_low_beta().rename("low_beta")
    msci = load_msci()
    write_premia(factors)
    write_performance(msci, factors["RF"])
    write_corr(msci, factors, mom, low_beta)
    write_momentum_vol(factors["Mkt-RF"], mom)


if __name__ == "__main__":
    main()
