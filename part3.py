"""Part 3. Industry minimum-variance portfolios.

    python part3.py
"""

from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import LogLocator, ScalarFormatter
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from scipy.optimize import minimize

ROOT = Path(__file__).resolve().parent
IND = ROOT / "data" / "industries.xlsx"
FF = ROOT / "data" / "ff_msci.xlsx"
OUT = ROOT / "exhibits" / "part3"

plt.rcParams.update({
    "figure.dpi": 140,
    "savefig.bbox": "tight",
    "savefig.facecolor": "white",
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "axes.edgecolor": "#333333",
    "axes.labelcolor": "#222222",
    "axes.titlecolor": "#1F4E79",
    "text.color": "#222222",
    "axes.grid": True,
    "grid.color": "#E6E6E6",
    "grid.linewidth": 0.7,
    "axes.axisbelow": True,
    "font.size": 10,
})
BLUE, ORANGE = "#4C78A8", "#F58518"

NAVY = PatternFill("solid", fgColor="1F4E79")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
SUM = PatternFill("solid", fgColor="F2F2F2")


def month_stamp(value):
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        return pd.Period(pd.Timestamp(value), freq="M")
    except (ValueError, TypeError):
        return None


def as_period(value):
    if value is None or isinstance(value, datetime):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) == 6 and text.isdigit():
        return pd.Period(year=int(text[:4]), month=int(text[4:]), freq="M")
    return None


def take(frame, start, end):
    out = frame.loc[start:end]
    expect = (pd.Period(end, freq="M") - pd.Period(start, freq="M")).n + 1
    if len(out) != expect:
        raise RuntimeError(f"{start} to {end}: n={len(out)}, expected {expect}")
    return out


def load_industries():
    wb = load_workbook(IND, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    start = next(i for i, row in enumerate(rows) if as_period(row[0]) is not None)
    names = [str(cell).strip() for cell in rows[start - 1][1:] if cell is not None and str(cell).strip()]
    data = []
    index = []
    for row in rows[start:]:
        period = as_period(row[0])
        if period is None:
            break
        values = []
        for cell in row[1:1 + len(names)]:
            value = np.nan if cell is None else float(cell)
            values.append(np.nan if value <= -99 else value / 100.0)
        data.append(values)
        index.append(period)
    df = pd.DataFrame(data, index=index, columns=names).sort_index()
    return take(df, "1995-01", "2023-07")


def load_usa_levels():
    wb = load_workbook(FF, read_only=True, data_only=True)
    rows = list(wb["MSCI - factor totret index"].iter_rows(values_only=True))
    wb.close()
    header_at = next(i for i, row in enumerate(rows) if row and str(row[0]).strip() == "Date")
    usa_at = next(i for i, cell in enumerate(rows[header_at]) if cell and "USA" in str(cell).upper() and "VALUE" not in str(cell).upper())
    levels = {}
    for row in rows[header_at + 1:]:
        period = month_stamp(row[0] if row else None)
        if period is None:
            if levels:
                break
            continue
        levels[period] = float(str(row[usa_at]).replace(",", ""))
    return pd.Series(levels).sort_index()


def min_var(cov):
    ones = np.ones(cov.shape[0])
    raw = np.linalg.solve(cov, ones)
    return raw / raw.sum()


def min_var_long(cov):
    n = cov.shape[0]
    start = np.full(n, 1.0 / n)

    def objective(w):
        return w @ cov @ w

    def gradient(w):
        return 2.0 * cov @ w

    result = minimize(
        objective, start, jac=gradient, method="SLSQP",
        bounds=[(0.0, 1.0)] * n,
        constraints={"type": "eq", "fun": lambda w: w.sum() - 1.0},
        options={"ftol": 1e-12, "maxiter": 2000},
    )
    if not result.success:
        raise RuntimeError(result.message)
    weights = np.clip(result.x, 0.0, None)
    return weights / weights.sum()


def inv_vol(returns):
    scale = 1.0 / returns.std(ddof=1).to_numpy()
    return scale / scale.sum()


def realized(weights, returns):
    portfolio = returns.to_numpy() @ weights
    return float(pd.Series(portfolio).std(ddof=1) * np.sqrt(12)), portfolio


def write_weights(names, books):
    headers = ["Industry", "Min-var (shorts)", "Min-var (long-only)", "Inverse-vol"]
    columns = [books[key] for key in ("shorts", "long", "inv")]
    wb = Workbook()
    ws = wb.active
    ws.title = "weights"
    ws["A1"] = (
        "Industry weights as of 31 Dec 2004. Covariance from Jan 1995–Dec 2004. "
        "Green is the top 5 and red is the bottom 5 within each column. "
        "Ken French 49 value-weighted industry total returns."
    )
    ws["A1"].font = Font(italic=True, size=9, color="444444")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(2, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = NAVY
    for i, name in enumerate(names):
        ws.cell(i + 3, 1, name)
        for j, weights in enumerate(columns):
            ws.cell(i + 3, j + 2, float(weights[i]))
            ws.cell(i + 3, j + 2).number_format = "0.000000"
    sum_row = len(names) + 3
    ws.cell(sum_row, 1, "Sum").font = Font(bold=True)
    ws.cell(sum_row, 1).fill = SUM
    for j, weights in enumerate(columns):
        cell = ws.cell(sum_row, j + 2, float(weights.sum()))
        cell.number_format = "0.000000"
        cell.font = Font(bold=True)
        cell.fill = SUM
    for j, weights in enumerate(columns):
        order = np.argsort(weights, kind="mergesort")
        for rank in order[-5:]:
            ws.cell(int(rank) + 3, j + 2).fill = GREEN
        for rank in order[:5]:
            ws.cell(int(rank) + 3, j + 2).fill = RED
    ws.column_dimensions["A"].width = 36
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 22
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    wb.save(OUT / "3a_weights.xlsx")

    table = pd.DataFrame({"Industry": names, headers[1]: columns[0], headers[2]: columns[1], headers[3]: columns[2]})
    total = {"Industry": "Sum"}
    for header, weights in zip(headers[1:], columns):
        total[header] = float(weights.sum())
    table = pd.concat([table, pd.DataFrame([total])], ignore_index=True)
    table.to_csv(OUT / "3a_weights.csv", index=False, float_format="%.6f")


def write_vol(ind, books, long_2015, missing):
    blocks = [
        ("In-sample Jan 1995–Dec 2004, weights as of Dec 2004", take(ind, "1995-01", "2004-12"),
         (("Min-var (shorts)", books["shorts"]), ("Min-var (long-only)", books["long"]), ("Inverse-vol", books["inv"]))),
        ("Out-of-sample Jan 2005–Jul 2023, weights frozen at Dec 2004", take(ind, "2005-01", "2023-07"),
         (("Min-var (shorts)", books["shorts"]), ("Min-var (long-only)", books["long"]), ("Inverse-vol", books["inv"]))),
        ("Jan 2016–Jul 2023, long-only minimum variance", take(ind, "2016-01", "2023-07"),
         (("Weights as of Dec 2004", books["long"]), ("Weights as of Dec 2015", long_2015))),
    ]
    lines = [
        "Exhibit 3B–3D — annualized realized volatility",
        "49 Ken French value-weighted industry total returns. Volatility = std of monthly portfolio returns × sqrt(12), in percent.",
        f"Missing codes (−99.99 or −999) in Jan 1995–Jul 2023: {missing}",
        "",
    ]
    for title, returns, specs in blocks:
        lines.append(title)
        for label, weights in specs:
            vol, _ = realized(weights, returns)
            lines.append(f"  {label:<28}{100 * vol:8.2f}")
        lines.append("")
    (OUT / "3b_c_d.txt").write_text("\n".join(lines).rstrip() + "\n")


def write_wealth(ind, weights, levels):
    _, monthly = realized(weights, take(ind, "2005-01", "2023-07"))
    port = pd.Series(monthly, index=take(ind, "2005-01", "2023-07").index)
    port_wealth = pd.concat([pd.Series({pd.Period("2004-12"): 1.0}), (1.0 + port).cumprod()])
    base = float(levels.loc[pd.Period("2004-12")])
    mkt = (levels.loc[pd.Period("2004-12"):pd.Period("2023-07")] / base)
    dates_p = port_wealth.index.to_timestamp("M")
    dates_m = mkt.index.to_timestamp("M")

    fig, ax = plt.subplots(figsize=(8.2, 4.6))
    ax.axvspan(datetime(2008, 9, 1), datetime(2009, 2, 28), color=BLUE, alpha=0.12, zorder=0)
    ax.axvspan(datetime(2020, 2, 1), datetime(2020, 3, 31), color=ORANGE, alpha=0.18, zorder=0)
    ax.plot(dates_p, port_wealth.to_numpy(), color=BLUE, lw=1.6, label="Long-only min-var", zorder=3)
    ax.plot(dates_m, mkt.to_numpy(), color=ORANGE, lw=1.5, label="MSCI USA", zorder=3)
    ax.set_yscale("log", base=2)
    ax.yaxis.set_major_locator(LogLocator(base=2))
    ax.yaxis.set_major_formatter(ScalarFormatter())
    ax.set_ylabel("Value of $1")
    ax.set_title("Growth of $1, 31 Dec 2004–Jul 2023")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(OUT / "3e_wealth.png")
    plt.close(fig)


def write_crises(ind):
    windows = (
        ("Sep 2008–Feb 2009", "2008-09", "2009-02"),
        ("Feb 2020–Mar 2020", "2020-02", "2020-03"),
    )
    lines = [
        "Exhibit 3E — ten industries with the highest cumulative total return in each crisis",
        "Ken French 49 value-weighted industry portfolios. Cumulative return is the compounded monthly total return, in percent.",
        "",
    ]
    for label, start, end in windows:
        block = take(ind, start, end)
        cum = (1.0 + block).prod() - 1.0
        top = cum.sort_values(ascending=False).head(10)
        lines.append(label)
        for name, value in top.items():
            lines.append(f"  {name:<12}{100 * value:8.2f}")
        lines.append("")
    (OUT / "3e_crises.txt").write_text("\n".join(lines).rstrip() + "\n")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    ind = load_industries()
    if ind.shape[1] != 49:
        raise RuntimeError(f"expected 49 industries, got {ind.shape[1]}")
    missing = int(ind.isna().sum().sum())
    if missing:
        raise RuntimeError(f"{missing} missing industry returns in the assignment window")
    seed = take(ind, "1995-01", "2004-12")
    cov = np.cov(seed.to_numpy(), rowvar=False, ddof=1)
    books = {"shorts": min_var(cov), "long": min_var_long(cov), "inv": inv_vol(seed)}
    later = take(ind, "2006-01", "2015-12")
    long_2015 = min_var_long(np.cov(later.to_numpy(), rowvar=False, ddof=1))
    write_weights(list(ind.columns), books)
    write_vol(ind, books, long_2015, missing)
    write_wealth(ind, books["long"], load_usa_levels())
    write_crises(ind)


if __name__ == "__main__":
    main()
